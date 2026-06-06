from __future__ import annotations

import io
from datetime import datetime, timezone

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse
from pydantic import BaseModel
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.database import get_db
from app.models.user import User
from app.models.resume import ResumeVersion
from app.services.auth import get_current_user
from app.services import gemini
from app.services.ats_scorer import compute_ats_score
from app.config import get_settings

settings = get_settings()
router = APIRouter(tags=["builder"])


class BuildRequest(BaseModel):
    target_role: str = "Software Engineer"
    target_company: str | None = None
    user_data: dict = {}


@router.post("/resume/build")
async def build_resume(
    body: BuildRequest,
    user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    if not settings.GEMINI_API_KEY:
        content = body.user_data or _mock_resume_content()
    else:
        try:
            content = await gemini.build_resume(
                body.user_data, body.target_role, body.target_company
            )
        except Exception as e:
            raise HTTPException(500, f"Resume generation failed: {e}")

    text_repr = _content_to_text(content)
    ats = compute_ats_score(text_repr)

    version = ResumeVersion(
        user_id=user.id,
        version_name=f"{body.target_role} - {datetime.now(timezone.utc).strftime('%Y-%m-%d')}",
        content=content,
        ats_score=ats.composite_score,
        status="active",
    )
    db.add(version)
    await db.flush()

    return {
        "status": "success",
        "resume": {
            "id": str(version.id),
            "content": content,
            "ats_score": ats.composite_score,
            "version_name": version.version_name,
        },
    }


@router.get("/resume/download/{version_id}")
async def download_resume(
    version_id: str,
    format: str = "pdf",
    user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    result = await db.execute(
        select(ResumeVersion).where(
            ResumeVersion.id == version_id,
            ResumeVersion.user_id == user.id,
        )
    )
    version = result.scalar_one_or_none()
    if not version:
        raise HTTPException(404, "Resume version not found")

    if format == "docx":
        return _generate_docx(version.content, version.version_name)
    return _generate_pdf(version.content, version.version_name)


@router.get("/resume/versions")
async def list_versions(
    user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    result = await db.execute(
        select(ResumeVersion)
        .where(ResumeVersion.user_id == user.id)
        .order_by(ResumeVersion.created_at.desc())
    )
    versions = result.scalars().all()

    return {
        "status": "success",
        "versions": [
            {
                "id": str(v.id),
                "version_name": v.version_name,
                "ats_score": v.ats_score,
                "status": v.status,
                "created_at": v.created_at.isoformat(),
            }
            for v in versions
        ],
    }


def _content_to_text(content: dict) -> str:
    parts = []
    if content.get("name"):
        parts.append(content["name"])
    if content.get("title"):
        parts.append(content["title"])
    if content.get("summary"):
        parts.append(f"\nSummary\n{content['summary']}")
    for exp in content.get("experience", []):
        parts.append(f"\nExperience\n{exp.get('role', '')} at {exp.get('company', '')}")
        for b in exp.get("bullets", []):
            parts.append(f"• {b}")
    edu = content.get("education", {})
    if edu:
        parts.append(f"\nEducation\n{edu.get('degree', '')} - {edu.get('school', '')}")
    skills = content.get("skills", [])
    if skills:
        parts.append(f"\nSkills\n{', '.join(skills)}")
    return "\n".join(parts)


def _generate_pdf(content: dict, filename: str) -> StreamingResponse:
    from reportlab.lib.pagesizes import letter
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
    from reportlab.lib.styles import getSampleStyleSheet

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter, topMargin=50, bottomMargin=50)
    styles = getSampleStyleSheet()
    story = []

    if content.get("name"):
        story.append(Paragraph(content["name"], styles["Title"]))
    if content.get("title"):
        story.append(Paragraph(content["title"], styles["Heading2"]))
    if content.get("summary"):
        story.append(Spacer(1, 12))
        story.append(Paragraph("Professional Summary", styles["Heading3"]))
        story.append(Paragraph(content["summary"], styles["Normal"]))

    for exp in content.get("experience", []):
        story.append(Spacer(1, 12))
        story.append(Paragraph(
            f"{exp.get('role', '')} at {exp.get('company', '')} ({exp.get('period', '')})",
            styles["Heading4"],
        ))
        for b in exp.get("bullets", []):
            story.append(Paragraph(f"&bull; {b}", styles["Normal"]))

    edu = content.get("education", {})
    if edu:
        story.append(Spacer(1, 12))
        story.append(Paragraph("Education", styles["Heading3"]))
        story.append(Paragraph(
            f"{edu.get('degree', '')} - {edu.get('school', '')} ({edu.get('year', '')})",
            styles["Normal"],
        ))

    skills = content.get("skills", [])
    if skills:
        story.append(Spacer(1, 12))
        story.append(Paragraph("Skills", styles["Heading3"]))
        story.append(Paragraph(", ".join(skills), styles["Normal"]))

    doc.build(story)
    buffer.seek(0)

    return StreamingResponse(
        buffer,
        media_type="application/pdf",
        headers={"Content-Disposition": f'attachment; filename="{filename}.pdf"'},
    )


def _generate_docx(content: dict, filename: str) -> StreamingResponse:
    from docx import Document

    doc = Document()
    if content.get("name"):
        doc.add_heading(content["name"], level=0)
    if content.get("title"):
        doc.add_heading(content["title"], level=2)
    if content.get("summary"):
        doc.add_heading("Professional Summary", level=3)
        doc.add_paragraph(content["summary"])

    for exp in content.get("experience", []):
        doc.add_heading(
            f"{exp.get('role', '')} at {exp.get('company', '')} ({exp.get('period', '')})",
            level=3,
        )
        for b in exp.get("bullets", []):
            doc.add_paragraph(b, style="List Bullet")

    edu = content.get("education", {})
    if edu:
        doc.add_heading("Education", level=3)
        doc.add_paragraph(
            f"{edu.get('degree', '')} - {edu.get('school', '')} ({edu.get('year', '')})"
        )

    skills = content.get("skills", [])
    if skills:
        doc.add_heading("Skills", level=3)
        doc.add_paragraph(", ".join(skills))

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)

    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f'attachment; filename="{filename}.docx"'},
    )


def _mock_resume_content() -> dict:
    return {
        "name": "Your Name",
        "title": "Software Engineer",
        "summary": "Experienced software engineer with expertise in building scalable applications.",
        "experience": [
            {
                "company": "Tech Company",
                "role": "Software Engineer",
                "period": "2022 - Present",
                "bullets": ["Developed scalable microservices", "Improved system performance by 40%"],
            }
        ],
        "education": {"school": "University", "degree": "B.S. Computer Science", "year": "2022"},
        "skills": ["Python", "JavaScript", "React", "Node.js", "PostgreSQL"],
    }
